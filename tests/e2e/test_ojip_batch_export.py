"""
E2E tests for the multi-curve batch export curve sheets (params_summary.xlsx).

Covers the AquaPen multi-OJIP background bug (OJIP_raw exported the values as
uploaded, while F0/FJ/FI/FM are background-subtracted) and the new
OJIP_to_zero / OJIP_to_max / OJIP_norm / OJIP_uncorrected sheets, for the
three multi-curve sources: AquaPen/FluorPen .txt, Handy PEA .HAN and
OJIP Imaging .xlsx.

The JS under test lives in the MC module of js_OJIP.js (_basisValues,
_getNormValues, buildCurveSheets) and startBatchExport().
"""

import io
import zipfile

import numpy as np
import openpyxl
import pytest

from conftest import FIXTURES

ALL_SHEETS = {"raw": True, "shifted_F0": True, "shifted_FM": True,
              "double_norm": True, "uncorrected": True}
N_CURVES = 12


def _run_multicurve(page, fluorometer, fixture, n=N_CURVES, timeout=120000):
    """Upload a multi-curve file, keep the first n curves, run the batch pass."""
    page.select_option("#fluorometer", fluorometer)
    page.locator("#ojip-files").set_input_files(str(FIXTURES / fixture))
    page.locator("#analyze-btn").click()
    page.wait_for_selector("#mc-selection-modal", state="visible", timeout=60000)
    page.evaluate(
        """(n) => {
            document.querySelectorAll('.mc-curve-cb').forEach((cb, i) => { cb.checked = i < n; });
            MC._updateSelCount();
        }""", n)
    page.locator("#mc-analyze-btn").click()
    page.wait_for_function(
        """(n) => Array.isArray(paramMatrix) && paramMatrix.length === n
                  && paramMatrix.every(r => r)""",
        arg=n, timeout=timeout)
    rows = page.evaluate("paramMatrix.filter(r => r && !r.error)")
    assert len(rows) == n, [r for r in page.evaluate("paramMatrix") if r.get("error")]
    return rows


def _sheets(page, which=ALL_SHEETS):
    out = page.evaluate(
        "(w) => MC.buildCurveSheets(paramMatrix.filter(r => r && !r.error), w)", which)
    return {s["name"]: s["aoa"] for s in out}


def _col(aoa, j):
    """(time_ms, values) for curve column j (1-based after time), blanks → NaN."""
    t = np.array([r[0] for r in aoa[1:]], dtype=float)
    v = np.array([r[j] if r[j] != "" else np.nan for r in aoa[1:]], dtype=float)
    return t, v


def _assert_invariants(rows, sheets):
    raw, norm, to0, tomax = (sheets["OJIP_raw"], sheets["OJIP_norm"],
                             sheets["OJIP_to_zero"], sheets["OJIP_to_max"])
    assert raw[0][0] == "time_ms"
    assert raw[0][1:] == [r["name"] for r in rows]
    for j, r in enumerate(rows, start=1):
        t, y = _col(raw, j)
        ok = np.isfinite(y)
        # The Excel cross-check: interpolate OJIP_raw at the reported timing.
        assert np.interp(r["FJ_time_user_ms"], t[ok], y[ok]) == pytest.approx(r["FJ"], rel=1e-6)
        assert np.interp(r["FI_time_user_ms"], t[ok], y[ok]) == pytest.approx(r["FI"], rel=1e-6)
        assert np.nanmax(y) == pytest.approx(r["FM"], rel=1e-9)
        _, dn = _col(norm, j)
        assert np.nanmax(dn) == pytest.approx(1.0, abs=1e-9)
        _, z = _col(to0, j)
        np.testing.assert_allclose(z, y - r["F0"], atol=1e-6)
        _, m = _col(tomax, j)
        assert np.nanmax(m) == pytest.approx(0.0, abs=1e-6)


class TestAquapenMultiOjip:

    def test_curve_sheets_on_analysis_basis(self, ojip_page):
        page = ojip_page
        rows = _run_multicurve(page, "Aquapen", "aquapen_example_multi_OJIP.txt")
        for r in rows:
            assert r["basis_drop"] == 1 and r["basis_offset"] > 0
        sheets = _sheets(page)
        assert set(sheets) == {"OJIP_raw", "OJIP_to_zero", "OJIP_to_max",
                               "OJIP_norm", "OJIP_uncorrected"}
        _assert_invariants(rows, sheets)
        # The dropped 11 µs background row is gone from the corrected sheets…
        assert sheets["OJIP_raw"][1][0] == pytest.approx(0.021)
        # …but kept, with the original values, in OJIP_uncorrected.
        unc = sheets["OJIP_uncorrected"]
        assert unc[1][0] == pytest.approx(0.011)
        for j, r in enumerate(rows, start=1):
            assert unc[1][j] == pytest.approx(r["basis_offset"])   # 11 µs value == Bckg
            assert unc[2][j] - sheets["OJIP_raw"][1][j] == pytest.approx(r["basis_offset"])

    def test_sheet_selection(self, ojip_page):
        page = ojip_page
        _run_multicurve(page, "Aquapen", "aquapen_example_multi_OJIP.txt", n=3)
        only = _sheets(page, {"raw": False, "shifted_F0": False, "shifted_FM": False,
                              "double_norm": True, "uncorrected": False})
        assert list(only) == ["OJIP_norm"]

    def test_modal_checkboxes_default_on(self, ojip_page):
        page = ojip_page
        _run_multicurve(page, "Aquapen", "aquapen_example_multi_OJIP.txt", n=3)
        choice = page.evaluate("_batchExportCurveSheetChoice()")
        assert choice == ALL_SHEETS
        page.evaluate("document.getElementById('be-sheet-shifted-fm').checked = false")
        assert page.evaluate("_batchExportCurveSheetChoice()")["shifted_FM"] is False

    def test_zip_export_contains_curve_sheets(self, ojip_page):
        """End-to-end: the downloaded ZIP's params_summary.xlsx passes the
        user's Excel cross-check on every curve."""
        page = ojip_page
        rows = _run_multicurve(page, "Aquapen", "aquapen_example_multi_OJIP.txt", n=5)
        page.evaluate("showBatchExportModal()")
        page.wait_for_selector("#batchExportModal", state="visible", timeout=5000)
        with page.expect_download(timeout=120000) as dl:
            page.locator("#be-start-btn").click()
        data = open(dl.value.path(), "rb").read()
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            wb = openpyxl.load_workbook(io.BytesIO(zf.read("params_summary.xlsx")), data_only=True)
        for name in ("Parameters", "OJIP_raw", "OJIP_to_zero", "OJIP_to_max",
                     "OJIP_norm", "OJIP_uncorrected"):
            assert name in wb.sheetnames, wb.sheetnames
        sheets = {n: [[("" if c is None else c) for c in row] for row in wb[n].iter_rows(values_only=True)]
                  for n in ("OJIP_raw", "OJIP_to_zero", "OJIP_to_max", "OJIP_norm")}
        _assert_invariants(rows, sheets)


class TestAggregateNormalisation:

    def test_double_norm_values_peak_at_one(self, ojip_page):
        """_getNormValues (used by aggregate/panels/compare) must normalise the
        background-corrected curve — before the fix AquaPen peaked at ~1.1."""
        page = ojip_page
        _run_multicurve(page, "Aquapen", "aquapen_example_multi_OJIP.txt", n=4)
        peaks = page.evaluate(
            """() => MC.buildCurveSheets(paramMatrix.filter(r => r && !r.error),
                                         {double_norm: true})[0].aoa.slice(1)
                     .reduce((mx, row) => row.slice(1).map((v, i) => Math.max(mx[i] ?? -Infinity, v === '' ? -Infinity : v)), [])""")
        for p in peaks:
            assert p == pytest.approx(1.0, abs=1e-9)


class TestOtherMultiCurveSources:

    def test_handypea(self, ojip_page):
        page = ojip_page
        rows = _run_multicurve(page, "HandyPEA", "handypea_example.HAN", n=5)
        for r in rows:
            assert r["basis_drop"] == 0 and r["basis_offset"] == 0
        sheets = _sheets(page)
        assert "OJIP_uncorrected" not in sheets      # nothing was corrected
        _assert_invariants(rows, sheets)

    def test_ojip_imaging(self, ojip_page):
        page = ojip_page
        rows = _run_multicurve(page, "OJIPImaging", "ojipimaging_example.xlsx", n=5)
        sheets = _sheets(page)
        assert "OJIP_uncorrected" not in sheets
        _assert_invariants(rows, sheets)
